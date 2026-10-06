import io

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz

from app.config import Settings, get_settings
from app.db.database import Base, get_db
from app.main import app

SAMPLE = (
    "Natural language processing enables computers to analyse academic documents. "
    "Tokenization, lemmatization and TF-IDF are classical techniques."
)


@pytest.fixture
def settings() -> Settings:
    return Settings(max_upload_mb=1, min_text_chars=20)


@pytest.fixture
def make_pdf():
    def _make(pages: list[str]) -> bytes:
        doc = fitz.open()
        for text in pages:
            page = doc.new_page()
            if text:
                page.insert_textbox(fitz.Rect(50, 50, 550, 780), text, fontsize=11)
        data = doc.tobytes()
        doc.close()
        return data

    return _make


@pytest.fixture
def make_docx():
    def _make(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
        doc = DocxDocument()
        for p in paragraphs:
            doc.add_paragraph(p)
        if table:
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    t.cell(r, c).text = value
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    return _make


def isolated_settings(tmp_path, **overrides) -> Settings:
    """Settings whose artifact paths point at an empty temp folder, so tests never depend on a trained model."""
    values = dict(max_upload_mb=1, min_text_chars=20,
                  reference_idf_path=tmp_path / "no_reference_idf.json",
                  model_path=tmp_path / "no_model.joblib",
                  metrics_path=tmp_path / "no_metrics.json")
    values.update(overrides)
    return Settings(**values)


@pytest.fixture
def client(tmp_path):
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    TestingSession = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_get_db():
        db = TestingSession()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_settings] = lambda: isolated_settings(tmp_path)
    yield TestClient(app)  # no 'with' -> lifespan (real DB/log files) is not started
    app.dependency_overrides.clear()


@pytest.fixture
def tiny_limit_client(client, tmp_path):
    app.dependency_overrides[get_settings] = lambda: isolated_settings(tmp_path, max_upload_mb=0.001)
    return client


@pytest.fixture(scope="session")
def trained_artifacts(tmp_path_factory):
    """Train a small model once per test session in a temp folder (skipped without NLP resources)."""
    from app.nlp.resources import check_nlp_resources

    if not check_nlp_resources()["ready"]:
        pytest.skip("NLP resources not installed")
    from datasets.generate_dataset import generate
    from app.ml.train import train

    root = tmp_path_factory.mktemp("mini_dataset")
    generate(root / "raw", per_class=10, seed=1)
    generate(root / "challenge", per_class=4, seed=2)
    out = root / "artifacts"
    ref = root / "reference_idf.json"
    metrics = train(root / "raw", root / "challenge", root / "no_real_docs", out,
                    reference_idf_path=ref, cv_folds=3, seed=42)
    return {"root": root, "artifacts": out, "reference": ref, "metrics": metrics,
            "model": out / "model.joblib", "metrics_path": out / "metrics.json"}
